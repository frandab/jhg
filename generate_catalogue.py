from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
except ModuleNotFoundError:
    Workbook = load_workbook = None  # type: ignore

FILE_XLSX = Path("catalogo.xlsx")


@dataclass
class Order:
    id: int
    date: date
    title: str
    code: str
    qty: int
    price: float
    discount: float
    dist_cost: float
    order_type: str
    notes: str = ""
    production_cost: float = 0.0
    author_pct: float = 0.0


def load_or_create_workbook():
    if Workbook is None:
        raise ModuleNotFoundError("openpyxl is required to use this script")
    if FILE_XLSX.exists():
        wb = load_workbook(FILE_XLSX)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    return wb


def ensure_data_entry_sheet(wb):
    if "DataEntry" not in wb.sheetnames:
        ws = wb.create_sheet("DataEntry")
        ws.append([
            "ID ordine",
            "Data",
            "Titolo",
            "Codice titolo",
            "Quantità",
            "Prezzo unitario",
            "Sconto",
            "Costo distribuzione",
            "Tipo ordine",
            "Note",
        ])


def ensure_template_sheet(wb):
    if "TemplateTitolo" not in wb.sheetnames:
        ws = wb.create_sheet("TemplateTitolo")
        headers = [
            "Data ordine",
            "Quantità",
            "Prezzo lordo",
            "Prezzo netto",
            "Costo distribuzione",
            "Costo produzione",
            "Incasso netto",
            "% Autore",
            "Utile netto dopo split",
        ]
        ws.append(headers)
        ws.append([
            "Totali",
            "=SUM(B3:B1000)",
            "=SUM(C3:C1000)",
            "=SUM(D3:D1000)",
            "=SUM(E3:E1000)",
            "=SUM(F3:F1000)",
            "=SUM(G3:G1000)",
            "",
            "=SUM(I3:I1000)",
        ])
        ws.sheet_state = "hidden"


def ensure_recap_sheet(wb):
    current_year = date.today().year
    next_year = current_year + 1
    if "RECAP" not in wb.sheetnames:
        ws = wb.create_sheet("RECAP")
        ws.append([
            "Codice",
            "Titolo",
            "Tiratura",
            "Vendute",
            "Tot Costi",
            "Giacenza",
            f"Vendite {current_year}",
            f"Vendite {next_year}",
            "Prezzo copertina",
            "Vendite totali",
            "Prezzo netto medio",
            "Entrate (€)",
            "Costo unitario",
            "Costi",
            "Utile netto (€)",
            "UTILE NETTO DOPO SPLIT con autori",
        ])
    else:
        ws = wb["RECAP"]
        ws["G1"] = f"Vendite {current_year}"
        ws["H1"] = f"Vendite {next_year}"


def ensure_forecast_sheets(wb):
    current_year = date.today().year
    next_year = current_year + 1
    for year in (current_year, next_year):
        name = f"Previsioni {year}"
        if name not in wb.sheetnames:
            ws = wb.create_sheet(name)
            ws.append(["Codice", "Titolo", f"Previsione vendite {year}"])


def compute_forecasts(wb):
    current_year = date.today().year
    forecasts = {}
    relevant_sheets = [
        s
        for s in wb.sheetnames
        if s
        not in (
            "DataEntry",
            "TemplateTitolo",
            "RECAP",
        )
        and not s.startswith("Previsioni ")
    ]
    for code in relevant_sheets:
        sheet = wb[code]
        sold = 0
        first_date = None
        for row in sheet.iter_rows(min_row=3, values_only=True):
            row_date = row[0]
            qty = row[1] or 0
            if row_date:
                row_date = date.fromisoformat(row_date)
                if row_date.year == current_year:
                    sold += qty
                if first_date is None or row_date < first_date:
                    first_date = row_date
        forecast_current = sold * (12 / 9)
        if first_date and (date.today() - first_date) < timedelta(days=547):
            forecast_next = forecast_current * 0.65
        else:
            forecast_next = forecast_current
        forecasts[code] = (forecast_current, forecast_next, sheet.title)
    return forecasts


def update_forecast_sheets(wb, forecasts):
    current_year = date.today().year
    next_year = current_year + 1
    sheets = {
        current_year: wb[f"Previsioni {current_year}"],
        next_year: wb[f"Previsioni {next_year}"],
    }
    for ws in sheets.values():
        ws.delete_rows(2, ws.max_row)
    for code, (curr, nxt, title) in forecasts.items():
        sheets[current_year].append([code, title, curr])
        sheets[next_year].append([code, title, nxt])


def ensure_base_sheets(wb):
    ensure_data_entry_sheet(wb)
    ensure_template_sheet(wb)
    ensure_recap_sheet(wb)
    ensure_forecast_sheets(wb)


def save_workbook(wb):
    wb.save(FILE_XLSX)


def create_workbook():
    wb = load_or_create_workbook()
    ensure_base_sheets(wb)
    save_workbook(wb)


def process_order(order: Order):
    wb = load_or_create_workbook()
    ensure_base_sheets(wb)
    data_sheet = wb["DataEntry"]
    data_sheet.append([
        order.id,
        order.date.isoformat(),
        order.title,
        order.code,
        order.qty,
        order.price,
        order.discount,
        order.dist_cost,
        order.order_type,
        order.notes,
    ])
    if order.code not in wb.sheetnames:
        template = wb["TemplateTitolo"]
        title_sheet = wb.copy_worksheet(template)
        title_sheet.title = order.code
        title_sheet.sheet_state = "visible"
    sheet = wb[order.code]
    net_price = order.price * (1 - order.discount)
    incasso_netto = net_price * order.qty - order.dist_cost - order.production_cost
    utile_autore = incasso_netto * order.author_pct
    utile_netto = incasso_netto - utile_autore
    sheet.append([
        order.date.isoformat(),
        order.qty,
        order.price * order.qty,
        net_price * order.qty,
        order.dist_cost,
        order.production_cost,
        incasso_netto,
        order.author_pct,
        utile_netto,
    ])
    update_recap_sheet(wb)
    save_workbook(wb)


def update_recap_sheet(wb):
    forecasts = compute_forecasts(wb)
    update_forecast_sheets(wb, forecasts)
    recap = wb["RECAP"]
    recap.delete_rows(2, recap.max_row)
    relevant_sheets = [
        s
        for s in wb.sheetnames
        if s not in ("DataEntry", "TemplateTitolo", "RECAP") and not s.startswith("Previsioni ")
    ]
    for code in relevant_sheets:
        sheet = wb[code]
        vendute = sheet["B2"].value
        tot_costi = (sheet["E2"].value or 0) + (sheet["F2"].value or 0)
        incassi = sheet["G2"].value
        utile = sheet["I2"].value
        forecast_current, forecast_next, _ = forecasts.get(code, (None, None, None))
        recap.append([
            code,
            sheet.title,
            "",
            vendute,
            tot_costi,
            "",
            forecast_current,
            forecast_next,
            "",
            vendute,
            "",
            incassi,
            "",
            tot_costi,
            utile,
            utile,
        ])


if __name__ == "__main__":
    create_workbook()
